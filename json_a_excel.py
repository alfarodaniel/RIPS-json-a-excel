# -*- coding: utf-8 -*-
"""
Convierte archivos RIPS en JSON a Excel (un libro .xlsx por cada .json).

Uso: dejar este script (o el .exe generado con PyInstaller) en la carpeta que
contiene los .json y ejecutarlo. Por cada "X.json" se crea "X_convertido.xlsx".

Hojas generadas:
  Transaccion  -> encabezado de la factura + documento del usuario
  Usuarios     -> datos de cada usuario (numFactura al inicio)
  <servicio>   -> una hoja por cada tipo de servicio presente en el JSON
                  (Consultas, Medicamentos, Procedimientos, Hospitalizacion,
                   RecienNacidos, OtrosServicios, Urgencias, ...), con las
                  columnas de enlace numDocumentoIdentificacionUsuario,
                  numFactura y UstipoDocumentoIdentificacion al final.

Requisito: pip install openpyxl
"""
import json
import os
import sys
import traceback

from openpyxl import Workbook

# Nombre de hoja para cada clave de "servicios" (orden = orden en el JSON)
NOMBRES_HOJA = {
    "consultas": "Consultas",
    "medicamentos": "Medicamentos",
    "procedimientos": "Procedimientos",
    "urgencias": "Urgencias",
    "hospitalizacion": "Hospitalizacion",
    "recienNacidos": "RecienNacidos",
    "otrosServicios": "OtrosServicios",
}

CAMPOS_ENCABEZADO = ["numDocumentoIdObligado", "numFactura", "tipoNota", "numNota"]


def carpeta_base():
    """Carpeta donde está el .exe (o el .py)."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def cargar_json(ruta):
    with open(ruta, "r", encoding="utf-8-sig") as f:
        return json.load(f)


def celda(v):
    """Valores no escalares (listas/dicts) se guardan como texto JSON."""
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v


def escribir_hoja(wb, nombre, columnas, filas):
    ws = wb.create_sheet(nombre[:31])
    ws.append(columnas)
    for fila in filas:
        ws.append([celda(fila.get(c)) for c in columnas])
    return ws


def columnas_de(filas, base=()):
    """Unión ordenada de las claves de todas las filas."""
    cols = list(base)
    for f in filas:
        for k in f:
            if k not in cols:
                cols.append(k)
    return cols


def convertir(ruta_json, ruta_xlsx):
    data = cargar_json(ruta_json)
    usuarios = data.get("usuarios") or []
    num_factura = data.get("numFactura")

    wb = Workbook()
    wb.remove(wb.active)

    # --- Transaccion: una fila por usuario ---
    filas_tx = []
    for u in usuarios or [{}]:
        fila = {k: data.get(k) for k in CAMPOS_ENCABEZADO}
        fila["UstipoDocumentoIdentificacion"] = u.get("tipoDocumentoIdentificacion")
        fila["numDocumentoIdentificacion"] = u.get("numDocumentoIdentificacion")
        filas_tx.append(fila)
    cols_tx = CAMPOS_ENCABEZADO + ["UstipoDocumentoIdentificacion", "numDocumentoIdentificacion"]
    escribir_hoja(wb, "Transaccion", cols_tx, filas_tx)

    # --- Usuarios ---
    filas_us = []
    for u in usuarios:
        fila = {"numFactura": num_factura}
        fila.update({k: v for k, v in u.items() if k != "servicios"})
        filas_us.append(fila)
    cols_us = columnas_de(filas_us, ["numFactura"])
    escribir_hoja(wb, "Usuarios", cols_us, filas_us)

    # --- Servicios: se juntan los de todos los usuarios ---
    servicios = {}  # clave -> lista de filas (conserva orden de aparición)
    for u in usuarios:
        for clave, items in (u.get("servicios") or {}).items():
            for item in items or []:
                fila = dict(item)
                fila["numDocumentoIdentificacionUsuario"] = u.get("numDocumentoIdentificacion")
                fila["numFactura"] = num_factura
                fila["UstipoDocumentoIdentificacion"] = u.get("tipoDocumentoIdentificacion")
                servicios.setdefault(clave, []).append(fila)

    enlace = ["numDocumentoIdentificacionUsuario", "numFactura", "UstipoDocumentoIdentificacion"]
    for clave, filas in servicios.items():
        if not filas:
            continue
        nombre = NOMBRES_HOJA.get(clave, clave[:1].upper() + clave[1:])
        cols = [c for c in columnas_de(filas) if c not in enlace] + enlace
        escribir_hoja(wb, nombre, cols, filas)

    wb.save(ruta_xlsx)
    return {n: wb[n].max_row - 1 for n in wb.sheetnames}


def main():
    carpeta = carpeta_base()
    archivos = sorted(f for f in os.listdir(carpeta) if f.lower().endswith(".json"))
    if not archivos:
        print(f"No se encontraron archivos .json en: {carpeta}")
    ok = err = 0
    for nombre in archivos:
        ruta = os.path.join(carpeta, nombre)
        salida = os.path.join(carpeta, os.path.splitext(nombre)[0] + "_convertido.xlsx")
        try:
            resumen = convertir(ruta, salida)
            ok += 1
            print(f"[OK]    {nombre} -> {os.path.basename(salida)}")
            print("        " + ", ".join(f"{h}: {n}" for h, n in resumen.items()))
        except Exception as e:
            err += 1
            print(f"[ERROR] {nombre}: {e}")
            traceback.print_exc()
    print(f"\nTerminado. Convertidos: {ok}  Con error: {err}")
    if getattr(sys, "frozen", False):
        input("Presione Enter para cerrar...")


if __name__ == "__main__":
    main()
